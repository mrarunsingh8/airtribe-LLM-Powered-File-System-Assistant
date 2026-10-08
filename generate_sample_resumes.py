"""Generate 8 dummy resumes (PDF, DOCX, TXT) into ./resumes. Re-runnable."""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

RESUMES = {
    "resume_john_doe.pdf": """John Doe
john.doe@example.com | +1 555 0101 | Austin, TX
SUMMARY
Backend engineer with 6 years of Python experience building REST APIs and data pipelines.
EXPERIENCE
Senior Software Engineer, Acme Corp (2021-Present): Built Django and FastAPI services handling 2M requests/day; migrated jobs to Airflow.
Software Engineer, Initech (2018-2021): Wrote Python ETL scripts, PostgreSQL schemas and AWS Lambda functions.
SKILLS
Python, Django, FastAPI, PostgreSQL, AWS, Docker, Airflow
EDUCATION
B.S. Computer Science, University of Texas, 2018""",
    "resume_priya_sharma.pdf": """Priya Sharma
priya.sharma@example.com | +91 98765 43210 | Bengaluru, India
SUMMARY
Data scientist specialising in NLP and recommendation systems.
EXPERIENCE
Data Scientist, Flipcart Analytics (2020-Present): Trained transformer models in PyTorch; Python and SQL for feature engineering.
Analyst, DataWorks (2018-2020): Built dashboards in Tableau and forecasting models in R.
SKILLS
Python, PyTorch, scikit-learn, SQL, R, Tableau, Spark
EDUCATION
M.Tech Data Science, IIT Delhi, 2018""",
    "resume_maria_garcia.docx": """Maria Garcia
maria.garcia@example.com | +34 600 123 456 | Madrid, Spain
SUMMARY
Frontend developer focused on accessible, high-performance web apps.
EXPERIENCE
Frontend Lead, WebNova (2019-Present): Led React and TypeScript rewrite; improved Lighthouse scores by 40%.
UI Developer, PixelSoft (2016-2019): Built Angular components and design systems.
SKILLS
JavaScript, TypeScript, React, Next.js, CSS, Figma, Jest
EDUCATION
B.Eng Software Engineering, Universidad Politécnica de Madrid, 2016""",
    "resume_chen_wei.docx": """Chen Wei
chen.wei@example.com | +86 138 0000 1111 | Shanghai, China
SUMMARY
DevOps engineer automating cloud infrastructure at scale.
EXPERIENCE
DevOps Engineer, CloudPeak (2020-Present): Managed Kubernetes clusters with Terraform; wrote Python and Bash automation tooling.
Systems Administrator, NetCore (2017-2020): Linux administration, CI/CD with Jenkins.
SKILLS
Kubernetes, Terraform, AWS, GCP, Python, Bash, Jenkins, Prometheus
EDUCATION
B.S. Information Technology, Fudan University, 2017""",
    "resume_aisha_khan.txt": """Aisha Khan
aisha.khan@example.com | +44 7700 900123 | London, UK
SUMMARY
Product manager with a background in fintech and mobile payments.
EXPERIENCE
Senior Product Manager, PayFlow (2021-Present): Owned roadmap for mobile wallet with 3M users; ran A/B tests.
Product Analyst, BankLabs (2018-2021): Defined KPIs, wrote SQL reports, worked with engineering on Agile sprints.
SKILLS
Product strategy, Agile/Scrum, SQL, Jira, user research, A/B testing
EDUCATION
MBA, London Business School, 2018""",
    "resume_david_miller.txt": """David Miller
david.miller@example.com | +1 555 0199 | Seattle, WA
SUMMARY
Java engineer building low-latency trading systems.
EXPERIENCE
Staff Engineer, QuantEdge (2017-Present): Designed Java/Spring microservices and Kafka streaming pipelines.
Software Developer, FinSys (2013-2017): C++ and Java development for order-matching engines.
SKILLS
Java, Spring Boot, Kafka, C++, Redis, Microservices
EDUCATION
B.S. Computer Engineering, University of Washington, 2013""",
    "resume_fatima_ali.pdf": """Fatima Ali
fatima.ali@example.com | +971 50 123 4567 | Dubai, UAE
SUMMARY
Machine learning engineer deploying computer-vision models to production.
EXPERIENCE
ML Engineer, VisionAI (2021-Present): Built object-detection pipelines in Python with TensorFlow; served models via FastAPI on GCP.
Research Assistant, Khalifa University (2019-2021): Published 2 papers on medical image segmentation.
SKILLS
Python, TensorFlow, OpenCV, MLOps, Docker, GCP
EDUCATION
M.S. Artificial Intelligence, Khalifa University, 2021""",
    "resume_lucas_martin.docx": """Lucas Martin
lucas.martin@example.com | +33 6 12 34 56 78 | Paris, France
SUMMARY
UX/UI designer crafting intuitive B2B SaaS experiences.
EXPERIENCE
Lead Designer, SaaSly (2020-Present): Redesigned onboarding, raising activation by 25%; ran usability studies.
Designer, Studio Bleu (2017-2020): Brand identities and marketing websites.
SKILLS
Figma, Sketch, prototyping, user research, design systems, HTML/CSS
EDUCATION
Master of Design, École de Design Nantes, 2017""",
}


def main(out_dir: str = "resumes") -> None:
    out = Path(out_dir)
    out.mkdir(exist_ok=True)
    styles = getSampleStyleSheet()
    for name, text in RESUMES.items():
        lines = text.splitlines()
        path = out / name
        if name.endswith(".txt"):
            path.write_text(text + "\n", encoding="utf-8")
        elif name.endswith(".docx"):
            doc = Document()
            doc.add_heading(lines[0], level=0)
            for line in lines[1:]:
                if line.isupper():
                    doc.add_heading(line.title(), level=1)
                else:
                    doc.add_paragraph(line)
            doc.save(path)
        else:
            story = [Paragraph(lines[0], styles["Title"])]
            for line in lines[1:]:
                style = styles["Heading2"] if line.isupper() else styles["Normal"]
                story += [Paragraph(line, style), Spacer(1, 4)]
            SimpleDocTemplate(str(path), pagesize=A4).build(story)
        print(f"created {path}")


if __name__ == "__main__":
    main()
