"""
fs_tools.py — File-system tools designed to be called by an LLM.

Every tool returns plain JSON-serialisable data (dicts / lists) so results can
be passed straight back to the model as a tool result. Errors never raise out
of a tool; they are returned as ``{"success": False, "error": "..."}`` so the
LLM can read the failure and recover.

Supported read formats: .txt, .md, .pdf (pypdf), .docx (python-docx).
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

SUPPORTED_READ_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB safety cap


# --------------------------------------------------------------------------- #
# Internal helpers
# --------------------------------------------------------------------------- #
def _error(message: str, **extra: Any) -> dict:
    """Uniform error payload."""
    return {"success": False, "error": message, **extra}


def _iso(ts: float) -> str:
    return datetime.fromtimestamp(ts).isoformat(timespec="seconds")


def _normalise_ext(extension: Optional[str]) -> Optional[str]:
    if not extension:
        return None
    extension = extension.strip().lower()
    return extension if extension.startswith(".") else f".{extension}"


def _extract_pdf(path: Path) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages).strip(), {"pages": len(reader.pages)}


def _extract_docx(path: Path) -> tuple[str, dict]:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    # Include table cell text too — resumes often use tables for layout.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip(), {"paragraphs": len(document.paragraphs)}


def _extract_text(path: Path) -> tuple[str, dict]:
    for encoding in ("utf-8", "latin-1"):
        try:
            return path.read_text(encoding=encoding).strip(), {"encoding": encoding}
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode text file")


def _extract(path: Path) -> tuple[str, dict]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    return _extract_text(path)


# --------------------------------------------------------------------------- #
# Tool 1: read_file
# --------------------------------------------------------------------------- #
def read_file(filepath: str) -> dict:
    """Read a resume/document and return its text plus metadata.

    Args:
        filepath: Path to a .pdf, .docx, .txt or .md file.

    Returns:
        On success::

            {"success": True, "filepath": ..., "filename": ..., "extension": ...,
             "content": "<text>", "metadata": {"size_bytes", "modified",
             "word_count", "char_count", ...format-specific keys}}

        On failure: ``{"success": False, "error": "<reason>"}``
    """
    try:
        path = Path(filepath).expanduser()
        if not path.exists():
            return _error(f"File not found: {filepath}")
        if not path.is_file():
            return _error(f"Not a file: {filepath}")

        ext = path.suffix.lower()
        if ext not in SUPPORTED_READ_EXTENSIONS:
            return _error(
                f"Unsupported file type '{ext}'. "
                f"Supported: {sorted(SUPPORTED_READ_EXTENSIONS)}"
            )

        stat = path.stat()
        if stat.st_size > MAX_FILE_SIZE_BYTES:
            return _error(f"File too large ({stat.st_size} bytes)")

        content, fmt_meta = _extract(path)
        if not content:
            return _error("File contains no extractable text (it may be scanned/empty)",
                          filepath=str(path))

        return {
            "success": True,
            "filepath": str(path),
            "filename": path.name,
            "extension": ext,
            "content": content,
            "metadata": {
                "size_bytes": stat.st_size,
                "modified": _iso(stat.st_mtime),
                "word_count": len(content.split()),
                "char_count": len(content),
                **fmt_meta,
            },
        }
    except PermissionError:
        return _error(f"Permission denied: {filepath}")
    except Exception as exc:  # corrupt PDF/DOCX etc.
        return _error(f"Failed to read {filepath}: {type(exc).__name__}: {exc}")


# --------------------------------------------------------------------------- #
# Tool 2: list_files
# --------------------------------------------------------------------------- #
def list_files(directory: str, extension: Optional[str] = None) -> list:
    """List files in a directory (non-recursive), optionally filtered by extension.

    Args:
        directory: Folder to list.
        extension: Optional filter such as ".pdf" or "pdf" (case-insensitive).

    Returns:
        A list of dicts ``{"name", "path", "extension", "size_bytes", "modified"}``
        sorted by name. If the directory is invalid, a single-element list
        containing an error dict is returned so the LLM sees the reason.
    """
    try:
        folder = Path(directory).expanduser()
        if not folder.exists():
            return [_error(f"Directory not found: {directory}")]
        if not folder.is_dir():
            return [_error(f"Not a directory: {directory}")]

        wanted = _normalise_ext(extension)
        results = []
        for entry in sorted(folder.iterdir(), key=lambda p: p.name.lower()):
            if not entry.is_file() or entry.name.startswith("."):
                continue
            if wanted and entry.suffix.lower() != wanted:
                continue
            stat = entry.stat()
            results.append({
                "name": entry.name,
                "path": str(entry),
                "extension": entry.suffix.lower(),
                "size_bytes": stat.st_size,
                "modified": _iso(stat.st_mtime),
            })
        return results
    except PermissionError:
        return [_error(f"Permission denied: {directory}")]
    except Exception as exc:
        return [_error(f"Failed to list {directory}: {exc}")]


# --------------------------------------------------------------------------- #
# Tool 3: write_file
# --------------------------------------------------------------------------- #
def write_file(filepath: str, content: str) -> dict:
    """Write text content to a file, creating parent directories as needed.

    Overwrites an existing file.

    Returns:
        ``{"success": True, "filepath", "bytes_written", "created": bool}``
        or ``{"success": False, "error"}``.
    """
    try:
        if content is None:
            return _error("No content provided")
        path = Path(filepath).expanduser()
        if path.exists() and path.is_dir():
            return _error(f"Path is a directory: {filepath}")

        existed = path.exists()
        path.parent.mkdir(parents=True, exist_ok=True)
        data = str(content).encode("utf-8")
        path.write_bytes(data)
        return {
            "success": True,
            "filepath": str(path),
            "bytes_written": len(data),
            "created": not existed,
        }
    except PermissionError:
        return _error(f"Permission denied: {filepath}")
    except Exception as exc:
        return _error(f"Failed to write {filepath}: {exc}")


# --------------------------------------------------------------------------- #
# Tool 4: search_in_file
# --------------------------------------------------------------------------- #
def search_in_file(filepath: str, keyword: str, context_chars: int = 60) -> dict:
    """Case-insensitive keyword search with surrounding context.

    Args:
        filepath: File to search (any format ``read_file`` supports).
        keyword: Word or phrase to find.
        context_chars: Characters of context to include on each side.

    Returns:
        ``{"success": True, "filepath", "keyword", "match_count",
        "matches": [{"line", "position", "match", "context"}]}``
    """
    if not keyword or not keyword.strip():
        return _error("Keyword must be a non-empty string")

    doc = read_file(filepath)
    if not doc.get("success"):
        return doc

    text = doc["content"]
    pattern = re.compile(re.escape(keyword.strip()), re.IGNORECASE)
    matches = []
    for m in pattern.finditer(text):
        start, end = m.start(), m.end()
        lo, hi = max(0, start - context_chars), min(len(text), end + context_chars)
        snippet = re.sub(r"\s+", " ", text[lo:hi]).strip()
        matches.append({
            "line": text.count("\n", 0, start) + 1,
            "position": start,
            "match": m.group(0),
            "context": ("…" if lo > 0 else "") + snippet + ("…" if hi < len(text) else ""),
        })

    return {
        "success": True,
        "filepath": doc["filepath"],
        "filename": doc["filename"],
        "keyword": keyword,
        "match_count": len(matches),
        "matches": matches,
    }


# --------------------------------------------------------------------------- #
# Tool registry — JSON-schema definitions consumed by the LLM layer
# --------------------------------------------------------------------------- #
TOOL_FUNCTIONS = {
    "read_file": read_file,
    "list_files": list_files,
    "write_file": write_file,
    "search_in_file": search_in_file,
}

TOOL_SCHEMAS = [
    {
        "name": "read_file",
        "description": "Read a document (PDF, DOCX, TXT, MD) and return its full text "
                       "content and metadata. Use this to look at a resume's contents.",
        "parameters": {
            "type": "object",
            "properties": {
                "filepath": {"type": "string", "description": "Path to the file, e.g. 'resumes/jane.pdf'"},
            },
            "required": ["filepath"],
        },
    },
    {
        "name": "list_files",
        "description": "List files in a directory with name, path, size and modified date. "
                       "Optionally filter by extension like '.pdf'.",
        "parameters": {
            "type": "object",
            "properties": {
                "directory": {"type": "string", "description": "Directory to list, e.g. 'resumes'"},
                "extension": {"type": "string", "description": "Optional extension filter, e.g. '.pdf'"},
            },
            "required": ["directory"],
        },
    },
    {
        "name": "write_file",
        "description": "Write text content to a file, creating parent folders if needed. "
                       "Overwrites existing files.",
        "parameters": {
            "type": "object",
            "properties": {
                "filepath": {"type": "string", "description": "Destination path, e.g. 'summaries/john.md'"},
                "content": {"type": "string", "description": "Full text to write"},
            },
            "required": ["filepath", "content"],
        },
    },
    {
        "name": "search_in_file",
        "description": "Case-insensitive search for a keyword in one file; returns each match "
                       "with surrounding context. Call once per file to search a folder.",
        "parameters": {
            "type": "object",
            "properties": {
                "filepath": {"type": "string", "description": "File to search"},
                "keyword": {"type": "string", "description": "Word or phrase to find"},
            },
            "required": ["filepath", "keyword"],
        },
    },
]


def execute_tool(name: str, arguments: dict) -> Any:
    """Dispatch a tool call by name. Unknown tools / bad args return an error dict."""
    func = TOOL_FUNCTIONS.get(name)
    if func is None:
        return _error(f"Unknown tool: {name}")
    try:
        return func(**(arguments or {}))
    except TypeError as exc:
        return _error(f"Invalid arguments for {name}: {exc}")
